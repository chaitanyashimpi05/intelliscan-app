import streamlit as st
import pdfplumber
import docx
import pandas as pd
import plotly.graph_objects as go
import os
from fpdf import FPDF
import spacy

# Load spaCy model (already installed via requirements.txt)
nlp = spacy.load("en_core_web_sm")

# --------- FUNCTIONS ---------

def extract_text(uploaded_file):
    file_extension = os.path.splitext(uploaded_file.name)[-1].lower()
    if file_extension == ".pdf":
        with pdfplumber.open(uploaded_file) as pdf:
            return "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
    elif file_extension == ".docx":
        doc = docx.Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    return None

predefined_skills = {
    "Software Engineer": {"Python", "Java", "C++", "Machine Learning", "Git", "SQL", "Docker", "Kubernetes", "JavaScript", "React", "Node.js"},
    "Data Scientist": {"Pandas", "NumPy", "Deep Learning", "Statistics", "Python", "R", "TensorFlow", "PyTorch", "SQL", "Data Visualization"},
    "Marketing": {"SEO", "Google Ads", "Content Writing", "Social Media", "Branding", "Google Analytics", "Copywriting", "Email Marketing"},
    "UI/UX Designer": {"Figma", "Adobe XD", "Wireframing", "User Research", "Sketch", "Prototyping", "Design Thinking"},
    "Cybersecurity Analyst": {"Network Security", "Ethical Hacking", "Cryptography", "Firewalls", "Incident Response", "Penetration Testing"}
}

def extract_skills_nlp(text):
    doc = nlp(text)
    return {token.text for token in doc if token.pos_ in ["NOUN", "PROPN"]}

def match_skills(resume_text, job_role):
    extracted_skills = extract_skills_nlp(resume_text)
    job_skills = predefined_skills.get(job_role, set())
    matched_skills = extracted_skills.intersection(job_skills)
    missing_skills = job_skills - matched_skills
    match_percentage = (len(matched_skills) / len(job_skills)) * 100 if job_skills else 0
    return matched_skills, missing_skills, match_percentage

def job_match_percentage(resume_text, job_description):
    resume_words = set(resume_text.lower().split())
    job_words = set(job_description.lower().split())
    match_count = len(resume_words.intersection(job_words))
    return (match_count / len(job_words)) * 100 if job_words else 0

def check_resume_format(resume_text):
    sections = {
        "Education": ["education", "degree", "bachelor", "master", "university"],
        "Experience": ["experience", "internship", "company", "work"],
        "Skills": ["skills", "technologies", "expertise"],
        "Certifications": ["certification", "course", "training"],
        "Projects": ["projects", "portfolio", "github"]
    }
    return [section for section, keywords in sections.items() if not any(keyword in resume_text.lower() for keyword in keywords)]

def generate_pdf(job_role, skill_match_score, job_match_score, matched_skills, missing_skills, missing_sections):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(200, 10, "IntelliScan Resume Analysis Report", ln=True, align="C")

    pdf.set_font("Arial", "", 12)
    pdf.cell(200, 10, f"Job Role: {job_role}", ln=True)
    pdf.cell(200, 10, f"Skill Match Score: {skill_match_score:.2f}%", ln=True)
    pdf.cell(200, 10, f"Job Match Score: {job_match_score:.2f}%", ln=True)

    pdf.ln(5)
    pdf.cell(200, 10, "Matched Skills:", ln=True)
    pdf.multi_cell(0, 10, ", ".join(matched_skills) if matched_skills else "None")

    pdf.ln(5)
    pdf.cell(200, 10, "Missing Skills:", ln=True)
    pdf.multi_cell(0, 10, ", ".join(missing_skills) if missing_skills else "None")

    pdf.ln(5)
    pdf.cell(200, 10, "Missing Resume Sections:", ln=True)
    pdf.multi_cell(0, 10, ", ".join(missing_sections) if missing_sections else "None")

    return pdf.output(dest="S").encode("latin1")

# --------- STREAMLIT UI ---------

st.set_page_config(page_title="IntelliScan - AI Resume Analyzer", layout="wide")
st.title("📄 IntelliScan - AI-Based Resume Analyzer with PDF Export")

with st.expander("📝 How It Works", expanded=False):
    st.markdown("""
    - Upload one or more **PDF or DOCX resumes**.
    - Paste the **job description**.
    - Select a **job role** from the dropdown.
    - Get insights on skill match, job relevance, formatting, and download a full **PDF report**.
    """)

uploaded_files = st.file_uploader("📤 Upload your resumes (PDF/DOCX)", type=["pdf", "docx"], accept_multiple_files=True)
job_description = st.text_area("🧾 Paste Job Description Here")
job_role = st.selectbox("💼 Select Job Role", list(predefined_skills.keys()))

if uploaded_files and job_description:
    for uploaded_file in uploaded_files:
        resume_text = extract_text(uploaded_file)
        if resume_text:
            matched_skills, missing_skills, skill_match_score = match_skills(resume_text, job_role)
            job_match_score = job_match_percentage(resume_text, job_description)
            missing_sections = check_resume_format(resume_text)

            st.divider()
            st.subheader(f"📊 Analysis for: `{uploaded_file.name}`")

            col1, col2 = st.columns(2)
            col1.metric("✅ Skill Match", f"{skill_match_score:.2f}%")
            col2.metric("🎯 Job Match", f"{job_match_score:.2f}%")

            st.markdown("### 🧠 Skills Overview")
            st.success(f"Matched Skills: {', '.join(matched_skills) if matched_skills else 'None'}")
            st.warning(f"Missing Skills: {', '.join(missing_skills) if missing_skills else 'None'}")

            st.markdown("### 📌 Resume Formatting Check")
            if missing_sections:
                st.error(f"❌ Missing Sections: {', '.join(missing_sections)}")
            else:
                st.success("✅ Resume is well-structured!")

            pdf_bytes = generate_pdf(job_role, skill_match_score, job_match_score, matched_skills, missing_skills, missing_sections)
            st.download_button(
                label=f"📥 Download Report for {uploaded_file.name}",
                data=pdf_bytes,
                file_name=f"{uploaded_file.name}_analysis.pdf",
                mime="application/pdf"
            )
